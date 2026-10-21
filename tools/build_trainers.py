#!/usr/bin/env python3
"""Genera web/data/trainers.json: tutti gli allenatori di Imperium con le loro squadre
e l'elenco ordinato degli scontri importanti di ogni parte della guida.

Fonti:
  - raw_sets.json  squadre dal calcolatore di hzla (tr_id = numero dell'allenatore nel gioco)
  - boss.xlsx      foglio "Boss Battles" della documentazione: ordine degli scontri per medaglia
Uso: python3 tools/build_trainers.py raw_sets.json boss.xlsx web/data/imperium.json web/data/trainers.json
"""
import json
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")
SETS, XLSX, DATA, OUT = sys.argv[1:5]
sets = json.load(open(SETS, encoding="utf8"))
data = json.load(open(DATA, encoding="utf8"))


def norm(s):
    return re.sub(r"[^a-z0-9]", "", s.lower())


# ---------- specie: nome del calcolatore -> id ----------
name_to_sid = {}
for sid, sp in data["species"].items():
    n = sp["n"]
    base, _, form = n.partition(" (")
    form = form.rstrip(")")
    name_to_sid.setdefault(norm(base + form), int(sid))
    name_to_sid.setdefault(norm(n), int(sid))
    if not form:
        name_to_sid.setdefault(norm(base), int(sid))


def sid_of(name):
    n = norm(name.replace("-Mega", "Mega"))
    if n in name_to_sid:
        return name_to_sid[n]
    base = norm(name.split("-")[0])
    return name_to_sid.get(base)


move_id = {norm(m[0]): int(k) for k, m in data["moves"].items()}

# ---------- allenatori ----------
trainers = {}
for sp_name, ts in sets.items():
    for label, st in ts.items():
        tid = st["tr_id"]
        t = trainers.setdefault(tid, {"n": re.sub(r"^Lvl -?\d+ ", "", label).strip(), "m": [], "d": False})
        t["m"].append((st.get("sub_index", 0), sp_name, st))
        if st.get("battle_type") == "Doubles":
            t["d"] = True

out_trainers = {}
for tid, t in trainers.items():
    mons = sorted(t["m"], key=lambda x: x[0])
    names = {m[1] for m in mons}
    packed = []
    for _, sp_name, st in mons:
        base = re.sub(r"-Mega(-[A-Z])?$|-Primal$", "", sp_name)
        if base != sp_name and base in names:
            for p in packed:
                if p["sp"] == base:
                    p["mega"] = sp_name
            continue
        lv = st.get("level")
        try:
            lv = int(lv)
        except (TypeError, ValueError):
            lv = None
        packed.append({"sp": sp_name, "sid": sid_of(sp_name), "lv": lv, "item": st.get("item"),
                       "ab": st.get("ability"), "nat": st.get("nature"),
                       "mv": [move_id.get(norm(x)) or x for x in st.get("moves", [])],
                       # IV ed EV nell'ordine PS, Att, Dif, AttSp, DifSp, Vel (servono a stimare statistiche e danni)
                       "iv": [st.get("ivs", {}).get(k, 31) for k in ("hp", "at", "df", "sa", "sd", "sp")],
                       "ev": [st.get("evs", {}).get(k, 0) for k in ("hp", "at", "df", "sa", "sd", "sp")]})
    out_trainers[tid] = {"n": t["n"], "d": t["d"], "m": [{k: v for k, v in p.items() if v not in (None, "", [])} for p in packed]}

# ---------- scontri importanti dal foglio ----------
SHEET_SPLIT = {"Badge 1 (Roxanne)": 0, "Badge 2 (Brawly)": 1, "Badge 3 (Wattson)": 2, "Badge 4 (Flannery)": 3,
               "Badge 5 (Norman)": 4, "Badge 6 (Winona)": 5, "Badge 7 (Tate & Liza)": 6, "Badge 8 (Juan)": 7,
               "Evergrande City": 8, "Victory Road": 8, "Elite 4  Champion": 8}
SINNOH = ["Roark", "Candice", "Gardenia", "Fantina", "Maylene", "Wake", "Byron", "Volkner"]
wb = openpyxl.load_workbook(XLSX, data_only=True)


def species_row(cells):
    """Riga con i nomi dei Pokémon della squadra (anche "(Lead Slot 1) / Plusle")."""
    names = [re.sub(r"^\(Lead Slot \d\)\s*/\s*", "", c).strip() for c in cells if c not in ("Rewards",)]
    names = [n for n in names if n and not n.startswith("(")]
    ok = [n for n in names if sid_of(n)]
    return ok if len(ok) >= max(1, len(names) - 1) else None


def title_ok(t):
    if t.startswith(("-", "TM", "Poke Paste", "Pokemon:", "Types:", "Level", "Speed", "The level cap", "(", "Gift Mon")):
        return False
    if re.search(r"ite[ XY]?$|Up:|Proc|Booster|Scarfed|Unburden|Bull Rush|Quill Rush|Sun Up|\d", t) and " - " not in t:
        return False
    return len(t) >= 4


fights = []
for ws in wb.worksheets:
    if ws.title not in SHEET_SPLIT:
        continue
    rows = [[str(c).strip() for c in r if c not in (None, "") and str(c).strip()] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if r]
    for k in range(len(rows) - 1):
        title = rows[k][0]
        if not title_ok(title) or sid_of(title):
            continue
        sp = species_row(rows[k + 1])
        if not sp:
            continue
        fights.append({"split": SHEET_SPLIT[ws.title], "title": title, "species": sp})


def match(fight):
    """Allenatori del calcolatore con la squadra più simile a quella del foglio."""
    want = {norm(s) for s in fight["species"]}
    head, _, place = fight["title"].partition(" - ")
    words = [w for w in re.split(r"[^A-Za-z]+", head) if len(w) > 2]
    where = [w for w in re.split(r"[^A-Za-z]+", place) if len(w) > 3 and w not in ("Optional", "Double", "Battle", "Permanent", "True")]
    scored = []
    for tid, t in out_trainers.items():
        have = {norm(m["sp"]) for m in t["m"]} | {norm(m.get("mega", "")) for m in t["m"]}
        overlap = len(want & have) / max(1, len(want))
        name_hit = sum(1 for w in words if w.lower() in t["n"].lower())
        place_hit = sum(1 for w in where if w.lower() in t["n"].lower())
        rematch = bool(re.search(r"\d$", t["n"]))  # Roxanne5 ecc. sono le rivincite
        if overlap >= 0.5:
            scored.append((overlap + 0.2 * name_hit + 0.1 * place_hit - (0.05 if rematch else 0), -tid, tid))
    scored.sort(reverse=True)
    if not scored:
        return []
    best = [scored[0][2]]
    base = lambda n: re.sub(r"\d+$", "", n)
    names = {base(out_trainers[best[0]]["n"])}
    for s, _, tid in scored[1:]:
        n = base(out_trainers[tid]["n"])
        if n in names:
            continue
        # varianti del rivale (stessa forza) o il secondo allenatore di una lotta doppia
        if "Rival" in fight["title"] and s >= scored[0][0] - 0.06:
            best.append(tid)
            names.add(n)
    # lotte doppie con due allenatori: il secondo si cerca per nome ("Drew & Branden", "Tate & Liza")
    m = re.search(r"&\s*(?:[A-Z][a-z]+\s)?([A-Z][a-z]+)", head)
    if m and not re.search(r"Grunt", m.group(0)):
        other = m.group(1)
        cands = sorted((bool(re.search(r"\d$", t["n"])), tid) for tid, t in out_trainers.items()
                       if re.search(r"\b" + other + r"\b", t["n"]) and tid not in best)
        if cands:
            best.append(cands[0][1])
    return best


# scontri che il foglio descrive in un formato diverso: aggiunti a mano
fights.insert(0, {"split": 0, "title": "Rivale - Route 103", "ids": [520, 523, 526, 529, 532, 535]})
fights.append({"split": 5, "title": "Aqua Admin Matt & Grunt - Weather Institute (Double Battle)", "ids": [18]})

bosses = []
seen = set()
for f in fights:
    ids = f.get("ids") or match(f)
    if not ids:
        print("senza squadra:", f["split"], f["title"], f["species"])
        continue
    key = (f["split"], tuple(ids))
    if key in seen:
        continue
    seen.add(key)
    title = re.sub(r"\s+", " ", f["title"]).strip()
    opt = bool(re.search(r"optional", title, re.I)) or any(re.search(r"\b" + s + r"\b", title) for s in SINNOH)
    bosses.append({"s": f["split"], "t": title, "ids": ids, "o": opt})

# ---------- dettagli in più dal foglio "Boss Battles" ----------
# Il foglio indica il tipo di Hidden Power e l'abilità delle forme Mega/Primal, che il calcolatore non ha.
EN_TYPES = {"normal": "Normale", "fighting": "Lotta", "flying": "Volante", "poison": "Veleno", "ground": "Terra", "rock": "Roccia",
            "bug": "Coleottero", "ghost": "Spettro", "steel": "Acciaio", "fire": "Fuoco", "water": "Acqua", "grass": "Erba",
            "electric": "Elettro", "psychic": "Psico", "ice": "Ghiaccio", "dragon": "Drago", "dark": "Buio", "fairy": "Folletto"}
sheet_mons = []
for ws in wb.worksheets:
    if ws.title not in SHEET_SPLIT and ws.title != "Sinnoh Gyms Only":
        continue
    rows = [[(str(c).strip() if c is not None else "") for c in r] for r in ws.iter_rows(values_only=True)]
    for i, r in enumerate(rows):
        if "Nature:" not in r:
            continue
        k = i
        while k > 0 and "Pokemon:" not in rows[k]:
            k -= 1
        sprow = rows[k - 3] if k >= 3 else []
        for cj in range(r.index("Nature:") + 1, len(r)):
            v = r[cj]
            if not v or v.startswith(("HP:", "Atk", "Def", "SpA", "SpD", "Spe")) or re.match(r"^[\d.]+$", v):
                continue
            g = lambda row: row[cj] if cj < len(row) else ""
            sp = (sprow[cj + 3] if cj + 3 < len(sprow) else "").split("/")[-1].split("\n")[-1].strip()
            mv = [rr[cj][1:].strip() for rr in rows[i + 3:i + 8] if cj < len(rr) and rr[cj].startswith("-")]
            sheet_mons.append({"sp": sp, "nat": v, "ab": g(rows[i + 1]), "mv": mv})

by_species = {}
for tid, t in out_trainers.items():
    for m in t["m"]:
        by_species.setdefault(norm(m["sp"]), []).append(m)
        if m.get("mega"):
            by_species.setdefault(norm(m["mega"]), []).append(m)
mvn = lambda x: data["moves"][str(x)][0] if isinstance(x, int) and str(x) in data["moves"] else str(x)
added_hp = added_ab = 0
for sm in sheet_mons:
    cands = by_species.get(norm(sm["sp"]), [])
    best = None
    for m in cands:
        have = {norm(mvn(x)) for x in m.get("mv", [])}
        sc = (norm(m.get("nat", "")) == norm(sm["nat"])) + len(have & {norm(x) for x in sm["mv"]})
        if not best or sc > best[0]:
            best = (sc, m)
    if not best or best[0] < 2:
        continue
    m = best[1]
    for x in sm["mv"]:
        hp = re.match(r"(?:HP|Hidden Power)\s*([A-Za-z]+)$", x)
        if hp and hp.group(1).lower() in EN_TYPES:
            m["hp"] = EN_TYPES[hp.group(1).lower()]
            added_hp += 1
    is_mega = norm(sm["sp"]) == norm(m.get("mega", ""))
    if is_mega and sm["ab"]:
        m["mab"] = sm["ab"].split("/")[0].strip()
        added_ab += 1
print("dal foglio: tipo di Hidden Power", added_hp, "· abilità Mega/Primal", added_ab)

json.dump({"trainers": out_trainers, "bosses": bosses}, open(OUT, "w", encoding="utf8"), ensure_ascii=False, separators=(",", ":"))
print(len(out_trainers), "allenatori,", len(bosses), "scontri importanti")
for b in bosses:
    if "Courtney & Grunt" in b["t"] and 146 not in b["ids"]:
        b["ids"].insert(0, 146)  # Magma Admin Courtney sul Mt. Chimney
bosses.sort(key=lambda b: b["s"])
for b in bosses:
    print(b["s"], "F" if b["o"] else "O", b["t"][:70], "->", [out_trainers[i]["n"] for i in b["ids"]][:3])
